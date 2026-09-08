"""Genera el Word de entrega EV02 (mismo contenido del PDF).

Requiere: pip install python-docx
Uso: python docs/ev02/generar_word.py
Sale: docs/ev02/EV02-InvenXis-GA8-220501096-AA1-EV02.docx
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

VERDE = RGBColor(0x00, 0x84, 0x3D)
APRENDIZ = '[NOMBRE DEL APRENDIZ]'
FICHA = '3235904'
INSTRUCTOR = 'Carlos Alberto Fuel Tulcán'
FECHA = '08 de septiembre de 2026'
EVIDENCIA = 'GA8-220501096-AA1-EV02'
REPO = 'https://github.com/Zultes-Dev/InvenXis'


def titulo(doc, texto, nivel=1):
    h = doc.add_heading(texto, level=nivel)
    for r in h.runs:
        r.font.color.rgb = VERDE
    return h


def tabla(doc, filas, anchos=None):
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, fila in enumerate(filas):
        for j, val in enumerate(fila):
            celda = t.cell(i, j)
            celda.text = ''
            p = celda.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if i == 0:
                r.bold = True
    if anchos:
        for j, a in enumerate(anchos):
            for i in range(len(filas)):
                t.cell(i, j).width = Cm(a)
    doc.add_paragraph('')
    return t


def main():
    doc = Document()
    for s in doc.sections:
        s.left_margin = Cm(2)
        s.right_margin = Cm(2)

    # PORTADA
    p = doc.add_paragraph()
    p.add_run('SERVICIO NACIONAL DE APRENDIZAJE — SENA\nCentro de Comercio y '
              'Turismo · Regional Quindío · ADSO').font.size = Pt(10)
    t = doc.add_paragraph()
    r = t.add_run('InvenXis — Sistema de Gestión de Inventarios')
    r.font.size = Pt(24)
    r.font.color.rgb = VERDE
    r.bold = True
    doc.add_heading('Evidencia GA8-220501096-AA1-EV02: Módulos Integrados', level=2)
    tabla(doc, [
        ['Campo', 'Valor'],
        ['Programa', 'Análisis y Desarrollo de Software (ADSO)'],
        ['Aprendiz', APRENDIZ],
        ['Ficha', FICHA],
        ['Instructor técnico', INSTRUCTOR],
        ['Fecha de entrega', FECHA],
        ['Repositorio', REPO],
        ['Versión', '1.0'],
    ], [4, 12])
    doc.add_paragraph(
        'Documento único de entrega: reúne los 5 ítems del instrumento de evaluación '
        '(módulos documentados, documento técnico, ambiente, control de versiones, acta '
        'de pruebas y aceptación) más manual de usuario y anexos. Se entrega en Zajuna '
        'junto con el enlace al repositorio.')
    doc.add_page_break()

    # CONTENIDO
    titulo(doc, 'Tabla de contenido')
    for t_ in ['Módulos codificados y documentados (ítem 1)',
                  'Documento técnico del sistema (ítem 2)',
                  'Ambiente de desarrollo y pruebas (ítem 3)',
                  'Código por control de versiones (ítem 4)',
                  'Acta de pruebas y aceptación (ítem 5)',
                  'Anexos']:
        doc.add_paragraph(f'{t_}', style='List Number')

    # ÍTEM 1
    titulo(doc, '1. Módulos codificados y documentados (ítem 1)')
    doc.add_paragraph('API base /api/ con envelope {"success": true, "data": ...} y JWT Bearer.')
    for nombre, entradas, salidas, archivos in [
        ('Autenticación y autorización',
         'POST /api/auth/login/, /api/auth/refresh/, /api/auth/logout/, GET /api/auth/me/',
         'JWT access+refresh con rotación y blacklist; throttle de login 20/min; RBAC por grupos.',
         'productos/api_views.py, config/settings.py, frontend/src/contexts/AuthContext.tsx'),
        ('Productos e inventario',
         'CRUD /api/productos/ + filtros y paginación',
         'Lee/escribe Producto (nombre, categoría, precio, cantidad, stock_min, proveedor).',
         'productos/models.py, api_views.py, frontend/src/pages/ProductosPage.tsx'),
        ('Proveedores y pedidos',
         'CRUD /api/proveedores/, /api/proveedores/{id}/pedidos/, /api/pedidos/{id}/',
         'Proveedor con conteos; pedidos con estados.',
         'productos/models.py, frontend/src/pages/ProveedoresPage.tsx'),
        ('Ventas (punto de venta)',
         'GET/POST /api/ventas/ con filtros desde/hasta',
         'Crea la venta y descuenta stock en transacción atómica; sin stock retorna 400 y revierte.',
         'productos/serializers.py:VentaCreateSerializer, frontend/src/pages/VentasPage.tsx'),
        ('Facturación y factura electrónica',
         'POST /api/ventas/{id}/facturar/, /api/facturas/, /emitir/, /anular/, /pdf/, /ubl/, /api/notas-credito/',
         'Consecutivo atómico FAC/NC, IVA 19%, CUFE SHA-384, UBL 2.1, proveedor DIAN intercambiable; anular crea NC.',
         'productos/fe/, productos/models.py:Factura, frontend/src/pages/FacturasPage.tsx'),
        ('Dashboard y reportes',
         'GET /api/dashboard/ (caché 60 s), /api/reportes/*, exportar excel|pdf',
         'KPIs, recientes, ranking, exportaciones Excel/PDF.',
         'productos/api_views.py, DashboardPage.tsx, ReportesPage.tsx'),
        ('Transversales',
         'GET /api/health/, /api/docs/, /api/redoc/, /admin/',
         'Health-check, OpenAPI, admin Django (9 modelos), throttling, caché.',
         'config/urls.py, productos/urls_api.py, productos/admin.py'),
    ]:
        titulo(doc, nombre, nivel=2)
        tabla(doc, [['Aspecto', 'Detalle'], ['Entradas', entradas],
                    ['Salidas / comportamiento', salidas], ['Código', archivos]], [3.5, 12.5])

    # ÍTEM 2
    titulo(doc, '2. Documento técnico del sistema (ítem 2)')
    doc.add_paragraph('2.1 Descripción general. InvenXis es un sistema de gestión de inventarios, '
                      'proveedores, ventas y facturación electrónica, con arquitectura API-first: backend '
                      'Django + DRF y SPA React.')
    doc.add_paragraph('2.2 Arquitectura. Petición HTTP → config/urls.py → productos/urls_api.py → '
                      'APIView/serializers → ORM → PostgreSQL/SQLite → JSON → TanStack Query → React. '
                      'Decisiones: transacciones atómicas, caché con invalidación, throttling, paginación '
                      'estándar y code-splitting por ruta.')
    doc.add_paragraph('2.3 Modelo de datos. Proveedor 1—N Producto; Proveedor 1—N Pedido 1—N DetallePedido; '
                      'Venta 1—N DetalleVenta; Venta 1—1 Factura 1—N NotaCredito; ContadorDocumento; User/Group '
                      'para RBAC.')
    titulo(doc, '2.4 Tecnologías y versiones (verificadas)', nivel=2)
    tabla(doc, [
        ['Capa', 'Tecnología y versión', 'Uso'],
        ['Lenguaje', 'Python 3.13.7', 'Backend'],
        ['Framework', 'Django 6.0.7', 'Servidor y ORM'],
        ['API', 'DRF 3.17.1 + SimpleJWT + drf-spectacular', 'REST, auth, OpenAPI'],
        ['Base de datos', 'SQLite (dev) / PostgreSQL 16 (prod)', 'Persistencia'],
        ['Caché', 'Redis 7 (prod) / local (dev)', 'Dashboard'],
        ['Servidor prod', 'Gunicorn + WhiteNoise (Docker)', 'Despliegue'],
        ['Lenguaje', 'TypeScript 6.0.3', 'Frontend tipado'],
        ['UI', 'React 19.2.8 + Vite 8.1.5 + Tailwind v4', 'SPA'],
        ['Datos', 'Axios + TanStack Query 5', 'HTTP, caché cliente'],
        ['Calidad', 'pytest + Playwright + oxlint + tsc', 'Tests y lint'],
        ['Runtime', 'Node 24.16.0 / npm 11.13.0', 'Frontend'],
    ], [3, 6, 7])
    doc.add_paragraph('2.5 Instalación y ejecución. Backend: python -m venv venv; pip install -r '
                      'requirements.txt; copiar .env.example a .env; python manage.py migrate; python '
                      'manage.py seed_data; python manage.py runserver (http://127.0.0.1:8000). Frontend: '
                      'cd frontend; npm install; npm run dev (http://localhost:3000). Demo: admin/admin, '
                      'operador/operador123. Producción: docker compose up --build.')
    doc.add_paragraph('2.6 Manual de usuario (resumen). Ingresar → Dashboard → Productos → Proveedores → '
                      'Ventas → Facturación (facturar, emitir FE, PDF/UBL, anular con nota crédito, solo '
                      'Administradores) → Reportes. El escudo del sidebar abre el admin Django.')
    doc.add_paragraph('2.7 Conclusiones. El sistema cumple el alcance con calidad verificable (45 pruebas '
                      'backend + 2 e2e en verde); queda como trabajo futuro la integración real con la DIAN.')

    # ÍTEM 3
    titulo(doc, '3. Ambiente de desarrollo y pruebas (ítem 3)')
    tabla(doc, [
        ['Herramienta', 'Versión'],
        ['Python', '3.13.7'],
        ['Django / DRF', '6.0.7 / 3.17.1'],
        ['Node.js / npm', '24.16.0 / 11.13.0'],
        ['React / Vite / TypeScript', '19.2.8 / 8.1.5 / 6.0.3'],
        ['Base de datos dev', 'SQLite (db.sqlite3, ignorada por git)'],
        ['Base de datos prod', 'PostgreSQL 16 + Redis 7 (docker-compose)'],
        ['Control de versiones', 'Git + GitHub'],
        ['Navegador de pruebas', 'Chromium (Playwright)'],
        ['SO de desarrollo', 'Windows + PowerShell 5.1'],
    ], [6, 10])
    doc.add_paragraph('Reproducción: 1) python -m venv venv y activar; 2) pip install -r requirements.txt; '
                      '3) copiar .env.example a .env; 4) python manage.py migrate && python manage.py seed_data; '
                      '5) python manage.py runserver; 6) cd frontend && npm install && npm run dev; '
                      '7) verificar con pytest -q y npm run build.')

    # ÍTEM 4
    titulo(doc, '4. Código por control de versiones (ítem 4)')
    doc.add_paragraph(f'Repositorio remoto (público): {REPO}, rama main, 18 commits convencionales '
                      '(feat/fix/chore/docs/ci). Historial reciente:')
    for c in ['ee2616f feat: modulo de ventas completo con facturacion y FE DIAN-ready',
              '6ca0e78 fix: login conserva next al admin + RBAC demo',
              '95e020e feat: acceso al panel Django admin desde el sidebar',
              '39e5ced feat: hardening senior x10 (health, throttle, factura uuid, a11y)',
              'cd0dabd feat: seed_data --force purga demo y resiembra determinista',
              '390eb1e feat: TanStack Query en SPA, API en /api, e2e Playwright, retiro SSR',
              '9d21c8c feat: P0+P1 hardening (venta atómica, logout blacklist, docker)',
              'd12bb07 chore: rename InvenSoft Pro to InvenXis']:
        doc.add_paragraph(c, style='List Bullet')
    doc.add_paragraph('CI (GitHub Actions): backend (pytest --cov), frontend (build + lint) y e2e '
                      '(Django + seed + Vite + Playwright).')

    # ÍTEM 5
    titulo(doc, '5. Acta de pruebas y aceptación (ítem 5)')
    doc.add_paragraph('5.1 Plan de pruebas manual (esperado vs. obtenido):')
    casos = [
        ('TC-01', 'Inicio de sesión válido', 'Acceso y panel', 'Acceso y panel (API + e2e)', 'Cumple'),
        ('TC-02', 'Crear producto', 'Registro creado', 'Creado vía API y UI', 'Cumple'),
        ('TC-03', 'Listar/buscar/paginar', 'Datos paginados', 'Paginación y filtros OK', 'Cumple'),
        ('TC-04', 'Editar producto', 'Dato actualizado', 'Actualizado vía API y UI', 'Cumple'),
        ('TC-05', 'Eliminar producto', 'Registro eliminado', 'Eliminado vía API y UI', 'Cumple'),
        ('TC-06', 'Registrar venta', 'Venta + descuento stock', 'Descuento atómico exacto', 'Cumple'),
        ('TC-07', 'Venta sin stock', 'Rechazo 400 sin venta', '400 + rollback total', 'Cumple'),
        ('TC-08', 'Facturar venta (IVA 19%)', 'Totales correctos', 'Subtotal/IVA/total exactos', 'Cumple'),
        ('TC-09', 'Emitir FE (CUFE+UBL)', 'Validada con CUFE/XML', 'Aceptada mock, XML válido', 'Cumple'),
        ('TC-10', 'Anular + RBAC', 'NC + stock; operador 403', 'NC + stock + 403 OK', 'Cumple'),
        ('TC-11', 'Logout + throttle', 'Blacklist; abuso 429', 'Blacklist y 429 OK', 'Cumple'),
        ('TC-12', 'DIAN en producción', 'Emisión real', 'Proveedor mock; falta certificado', 'No cumple'),
    ]
    tabla(doc, [['ID', 'Caso', 'Esperado', 'Obtenido', 'Estado']] + casos,
          [1.5, 3.5, 3.5, 4.5, 3])
    doc.add_paragraph('5.2 Ejecución y resultados. Backend: 45 pruebas pytest en verde, sin warnings. '
                      'Frontend: tsc + vite build OK, oxlint 0 errores. E2E Playwright: 2/2.')
    doc.add_paragraph('5.3 Observaciones y pendientes. TC-12 documentado con honestidad: la arquitectura '
                      'prevé el proveedor real (productos/fe/proveedores.py:get_proveedor), pero la emisión '
                      'productiva exige certificado digital, firma del UBL y rango autorizado. Prioridad media.')
    titulo(doc, '5.4 Acta de aceptación (formato)', nivel=2)
    tabla(doc, [['Campo', 'Valor'],
                ['Proyecto', 'InvenXis — Sistema de Gestión de Inventarios (v1.0)'],
                ['Evidencia', EVIDENCIA + ' — Módulos Integrados'],
                ['Aprendiz / Ficha', f'{APRENDIZ} / {FICHA}'],
                ['Fecha', FECHA],
                ['Resultado', '11 casos Cumplen / 1 pendiente documentado (TC-12)'],
                ['Soporte', f'Repositorio Git ({REPO}), pruebas automatizadas, manuales']],
          [4, 12])
    doc.add_paragraph('Declaración: el evaluador/cliente, luego de revisar la solución y verificar el '
                      'cumplimiento de los requisitos, acepta el entregable con el pendiente TC-12 documentado.')
    tabla(doc, [['Evaluador / Instructor', 'Aprendiz'],
                [f'Nombre: {INSTRUCTOR}\nFirma: ___________________________\nFecha: {FECHA}',
                 f'Nombre: {APRENDIZ}\nFirma: ___________________________\nFecha: {FECHA}']],
          [8, 8])

    # ANEXOS
    titulo(doc, 'A. Anexos')
    doc.add_paragraph('Endpoints (base /api/): auth/login, auth/refresh, auth/logout, auth/me, productos/, '
                      'proveedores/, pedidos, ventas, ventas/{id}/facturar/, facturas (+emitir/anular/pdf/ubl), '
                      'notas-credito/, dashboard/, reportes/*, exportar, health/, docs/, redoc/.')
    doc.add_paragraph('Credenciales demo (solo desarrollo): admin/admin, operador/operador123. App '
                      'http://localhost:3000, API http://127.0.0.1:8000, docs /api/docs/, admin /admin/, '
                      f'repo {REPO}.')

    out = 'docs/ev02/EV02-InvenXis-GA8-220501096-AA1-EV02.docx'
    doc.save(out)
    print(f'OK: {out}')


if __name__ == '__main__':
    main()
