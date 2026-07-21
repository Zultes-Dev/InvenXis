from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.paginator import Paginator
from django.shortcuts import render, redirect, get_object_or_404
from django.db import models
from django.db.models import Sum, F, Q
from .models import Producto, Proveedor
from .forms import ProductoForm, ProveedorForm

ITEMS_PER_PAGE = 10


def _get_base_context(request):
    return {
        'total_productos': Producto.objects.count(),
        'total_proveedores': Proveedor.objects.count(),
    }


@login_required
def dashboard(request):
    total_productos = Producto.objects.count()
    stock_bajo = Producto.objects.filter(cantidad__gt=0, cantidad__lte=models.F('stock_min')).count()
    sin_stock = Producto.objects.filter(cantidad=0).count()
    total_proveedores = Proveedor.objects.count()
    valor_total = Producto.objects.aggregate(total=Sum(F('cantidad') * F('precio')))['total'] or 0

    productos_recientes = Producto.objects.all()[:5]
    proveedores_activos = Proveedor.objects.filter(estado='activo')[:5]

    ctx = _get_base_context(request)
    ctx.update({
        'total_productos': total_productos,
        'stock_bajo': stock_bajo,
        'sin_stock': sin_stock,
        'total_proveedores': total_proveedores,
        'valor_total': valor_total,
        'productos_recientes': productos_recientes,
        'proveedores_activos': proveedores_activos,
    })
    return render(request, 'productos/dashboard.html', ctx)


@login_required
def lista_productos(request):
    filtro = request.GET.get('filtro', 'all')
    q = request.GET.get('q', '')

    productos = Producto.objects.all()

    if q:
        productos = productos.filter(
            Q(nombre__icontains=q) | Q(categoria__icontains=q)
        )

    if filtro == 'normal':
        productos = productos.filter(cantidad__gt=models.F('stock_min'))
    elif filtro == 'low':
        productos = productos.filter(cantidad__gt=0, cantidad__lte=models.F('stock_min'))
    elif filtro == 'critical':
        productos = productos.filter(cantidad=0)

    paginator = Paginator(productos, ITEMS_PER_PAGE)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    ctx = _get_base_context(request)
    ctx.update({
        'page_obj': page_obj,
        'filtro_actual': filtro,
        'query': q,
    })
    return render(request, 'productos/lista_productos.html', ctx)


@login_required
def detalle_producto(request, pk):
    producto = get_object_or_404(Producto, pk=pk)
    ctx = _get_base_context(request)
    ctx.update({'producto': producto})
    return render(request, 'productos/detalle_producto.html', ctx)


@login_required
def crear_producto(request):
    if request.method == 'POST':
        form = ProductoForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Producto creado correctamente.')
            return redirect('productos:lista')
    else:
        form = ProductoForm()

    ctx = _get_base_context(request)
    ctx.update({
        'form': form,
        'titulo': 'Registrar Nuevo Producto',
    })
    return render(request, 'productos/formulario_producto.html', ctx)


@login_required
def editar_producto(request, pk):
    producto = get_object_or_404(Producto, pk=pk)
    if request.method == 'POST':
        form = ProductoForm(request.POST, instance=producto)
        if form.is_valid():
            form.save()
            messages.success(request, 'Producto actualizado correctamente.')
            return redirect('productos:lista')
    else:
        form = ProductoForm(instance=producto)

    ctx = _get_base_context(request)
    ctx.update({
        'form': form,
        'titulo': 'Editar Producto',
        'producto': producto,
    })
    return render(request, 'productos/formulario_producto.html', ctx)


@login_required
def eliminar_producto(request, pk):
    producto = get_object_or_404(Producto, pk=pk)
    if request.method == 'POST':
        nombre = producto.nombre
        producto.delete()
        messages.success(request, f'Producto "{nombre}" eliminado correctamente.')
        return redirect('productos:lista')

    ctx = _get_base_context(request)
    ctx.update({'producto': producto})
    return render(request, 'productos/confirmar_eliminar.html', ctx)


@login_required
def lista_proveedores(request):
    filtro = request.GET.get('filtro', 'all')
    q = request.GET.get('q', '')

    proveedores = Proveedor.objects.all()

    if q:
        proveedores = proveedores.filter(
            Q(razon_social__icontains=q) | Q(nit__icontains=q) | Q(contacto__icontains=q)
        )

    if filtro == 'activo':
        proveedores = proveedores.filter(estado='activo')
    elif filtro == 'inactivo':
        proveedores = proveedores.filter(estado='inactivo')

    paginator = Paginator(proveedores, ITEMS_PER_PAGE)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    ctx = _get_base_context(request)
    ctx.update({
        'page_obj': page_obj,
        'filtro_actual': filtro,
        'query': q,
    })
    return render(request, 'productos/lista_proveedores.html', ctx)


@login_required
def crear_proveedor(request):
    if request.method == 'POST':
        form = ProveedorForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Proveedor creado correctamente.')
            return redirect('productos:lista_proveedores')
    else:
        form = ProveedorForm()

    ctx = _get_base_context(request)
    ctx.update({
        'form': form,
        'titulo': 'Registrar Nuevo Proveedor',
    })
    return render(request, 'productos/formulario_proveedor.html', ctx)


@login_required
def editar_proveedor(request, pk):
    proveedor = get_object_or_404(Proveedor, pk=pk)
    if request.method == 'POST':
        form = ProveedorForm(request.POST, instance=proveedor)
        if form.is_valid():
            form.save()
            messages.success(request, 'Proveedor actualizado correctamente.')
            return redirect('productos:lista_proveedores')
    else:
        form = ProveedorForm(instance=proveedor)

    ctx = _get_base_context(request)
    ctx.update({
        'form': form,
        'titulo': 'Editar Proveedor',
        'proveedor': proveedor,
    })
    return render(request, 'productos/formulario_proveedor.html', ctx)


@login_required
def eliminar_proveedor(request, pk):
    proveedor = get_object_or_404(Proveedor, pk=pk)
    if request.method == 'POST':
        nombre = proveedor.razon_social
        proveedor.delete()
        messages.success(request, f'Proveedor "{nombre}" eliminado correctamente.')
        return redirect('productos:lista_proveedores')

    ctx = _get_base_context(request)
    ctx.update({'proveedor': proveedor})
    return render(request, 'productos/confirmar_eliminar_proveedor.html', ctx)


@login_required
def reportes(request):
    total_productos = Producto.objects.count()
    stock_bajo = Producto.objects.filter(cantidad__gt=0, cantidad__lte=models.F('stock_min')).count()
    sin_stock = Producto.objects.filter(cantidad=0).count()
    stock_normal = Producto.objects.filter(cantidad__gt=models.F('stock_min')).count()
    total_proveedores = Proveedor.objects.count()
    proveedores_activos = Proveedor.objects.filter(estado='activo').count()
    proveedores_inactivos = Proveedor.objects.filter(estado='inactivo').count()
    valor_total = Producto.objects.aggregate(total=Sum(F('cantidad') * F('precio')))['total'] or 0

    ctx = _get_base_context(request)
    ctx.update({
        'total_productos': total_productos,
        'stock_bajo': stock_bajo,
        'sin_stock': sin_stock,
        'stock_normal': stock_normal,
        'total_proveedores': total_proveedores,
        'proveedores_activos': proveedores_activos,
        'proveedores_inactivos': proveedores_inactivos,
        'valor_total': valor_total,
    })
    return render(request, 'productos/reportes.html', ctx)


@login_required
def reporte_inventario(request):
    import json
    from django.db.models import Count, Sum as SumAgg
    from decimal import Decimal
    productos = Producto.objects.all()
    total_valor = productos.aggregate(total=Sum(F('cantidad') * F('precio')))['total'] or 0

    cats = productos.values('categoria').annotate(cnt=Count('id'), val=SumAgg(F('cantidad') * F('precio'))).order_by('categoria')
    categorias_labels = [c['categoria'] for c in cats]
    categorias_counts = [c['cnt'] for c in cats]
    categorias_valores = [float(c['val']) for c in cats]

    ctx = _get_base_context(request)
    ctx.update({
        'productos': productos,
        'total_productos': productos.count(),
        'stock_normal': productos.filter(cantidad__gt=models.F('stock_min')).count(),
        'stock_bajo': productos.filter(cantidad__gt=0, cantidad__lte=models.F('stock_min')).count(),
        'sin_stock': productos.filter(cantidad=0).count(),
        'total_valor': total_valor,
        'categorias_labels': json.dumps(categorias_labels),
        'categorias_counts': json.dumps(categorias_counts),
        'categorias_valores': json.dumps(categorias_valores),
        'reporte_tipo': 'inventario',
        'reporte_titulo': 'Reporte de Inventario',
    })
    return render(request, 'productos/reporte_detalle.html', ctx)


@login_required
def reporte_inventario_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Inventario'

    amber_fill = PatternFill(start_color='E8A230', end_color='E8A230', fill_type='solid')
    header_font = Font(name='Syne', bold=True, color='0D0E0F', size=11)
    data_font = Font(name='DM Mono', size=10, color='E8E6E1')
    thin_border = Border(
        left=Side(style='thin', color='252A2F'),
        right=Side(style='thin', color='252A2F'),
        top=Side(style='thin', color='252A2F'),
        bottom=Side(style='thin', color='252A2F'),
    )

    headers = ['Producto', 'Categoría', 'Stock', 'Stock Mínimo', 'Precio Unitario', 'Valor Total', 'Estado']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = amber_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    for i, p in enumerate(Producto.objects.all(), 2):
        estado = 'Sin Stock' if p.cantidad == 0 else ('Stock Bajo' if p.cantidad <= p.stock_min else 'Normal')
        ws.cell(row=i, column=1, value=p.nombre).font = data_font
        ws.cell(row=i, column=2, value=p.categoria).font = data_font
        ws.cell(row=i, column=3, value=p.cantidad).font = data_font
        ws.cell(row=i, column=4, value=p.stock_min).font = data_font
        ws.cell(row=i, column=5, value=float(p.precio)).font = data_font
        ws.cell(row=i, column=6, value=float(p.cantidad * p.precio)).font = data_font
        ws.cell(row=i, column=7, value=estado).font = data_font
        for col in range(1, 8):
            ws.cell(row=i, column=col).border = thin_border

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 15

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=reporte_inventario.xlsx'
    wb.save(response)
    return response


@login_required
def reporte_proveedores(request):
    import json
    from django.db.models import Count
    proveedores = Proveedor.objects.all()

    cats = proveedores.values('categoria').annotate(cnt=Count('id')).order_by('categoria')
    categorias_labels = [c['categoria'] for c in cats]
    categorias_counts = [c['cnt'] for c in cats]

    ctx = _get_base_context(request)
    ctx.update({
        'proveedores': proveedores,
        'total_proveedores': proveedores.count(),
        'proveedores_activos': proveedores.filter(estado='activo').count(),
        'proveedores_inactivos': proveedores.filter(estado='inactivo').count(),
        'categorias_labels': json.dumps(categorias_labels),
        'categorias_counts': json.dumps(categorias_counts),
        'reporte_tipo': 'proveedores',
        'reporte_titulo': 'Reporte de Proveedores',
    })
    return render(request, 'productos/reporte_detalle.html', ctx)


@login_required
def reporte_proveedores_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Proveedores'

    amber_fill = PatternFill(start_color='E8A230', end_color='E8A230', fill_type='solid')
    header_font = Font(name='Syne', bold=True, color='0D0E0F', size=11)
    data_font = Font(name='DM Mono', size=10, color='E8E6E1')
    thin_border = Border(
        left=Side(style='thin', color='252A2F'),
        right=Side(style='thin', color='252A2F'),
        top=Side(style='thin', color='252A2F'),
        bottom=Side(style='thin', color='252A2F'),
    )

    headers = ['Razón Social', 'NIT', 'Categoría', 'Contacto', 'Teléfono', 'Email', 'Estado']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = amber_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    for i, p in enumerate(Proveedor.objects.all(), 2):
        ws.cell(row=i, column=1, value=p.razon_social).font = data_font
        ws.cell(row=i, column=2, value=p.nit).font = data_font
        ws.cell(row=i, column=3, value=p.categoria).font = data_font
        ws.cell(row=i, column=4, value=p.contacto).font = data_font
        ws.cell(row=i, column=5, value=p.telefono).font = data_font
        ws.cell(row=i, column=6, value=p.email).font = data_font
        ws.cell(row=i, column=7, value='Activo' if p.estado == 'activo' else 'Inactivo').font = data_font
        for col in range(1, 8):
            ws.cell(row=i, column=col).border = thin_border

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 30
    ws.column_dimensions['G'].width = 12

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=reporte_proveedores.xlsx'
    wb.save(response)
    return response


@login_required
def reporte_completo(request):
    import json
    from django.db.models import Count
    productos = Producto.objects.all()
    proveedores = Proveedor.objects.all()
    total_valor = productos.aggregate(total=Sum(F('cantidad') * F('precio')))['total'] or 0

    ctx = _get_base_context(request)
    ctx.update({
        'productos': productos,
        'proveedores': proveedores,
        'total_productos': productos.count(),
        'stock_normal': productos.filter(cantidad__gt=models.F('stock_min')).count(),
        'stock_bajo': productos.filter(cantidad__gt=0, cantidad__lte=models.F('stock_min')).count(),
        'sin_stock': productos.filter(cantidad=0).count(),
        'total_proveedores': proveedores.count(),
        'proveedores_activos': proveedores.filter(estado='activo').count(),
        'proveedores_inactivos': proveedores.filter(estado='inactivo').count(),
        'total_valor': total_valor,
        'reporte_tipo': 'completo',
        'reporte_titulo': 'Reporte Completo',
    })
    return render(request, 'productos/reporte_detalle.html', ctx)


@login_required
def reporte_completo_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse

    wb = openpyxl.Workbook()

    amber_fill = PatternFill(start_color='E8A230', end_color='E8A230', fill_type='solid')
    header_font = Font(name='Syne', bold=True, color='0D0E0F', size=11)
    data_font = Font(name='DM Mono', size=10, color='E8E6E1')
    thin_border = Border(
        left=Side(style='thin', color='252A2F'),
        right=Side(style='thin', color='252A2F'),
        top=Side(style='thin', color='252A2F'),
        bottom=Side(style='thin', color='252A2F'),
    )

    # Productos sheet
    ws1 = wb.active
    ws1.title = 'Inventario'
    headers_p = ['Producto', 'Categoría', 'Stock', 'Stock Mínimo', 'Precio Unitario', 'Valor Total', 'Estado']
    for col, h in enumerate(headers_p, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.fill = amber_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
    for i, p in enumerate(Producto.objects.all(), 2):
        estado = 'Sin Stock' if p.cantidad == 0 else ('Stock Bajo' if p.cantidad <= p.stock_min else 'Normal')
        vals = [p.nombre, p.categoria, p.cantidad, p.stock_min, float(p.precio), float(p.cantidad * p.precio), estado]
        for col, v in enumerate(vals, 1):
            c = ws1.cell(row=i, column=col, value=v)
            c.font = data_font
            c.border = thin_border
    ws1.column_dimensions['A'].width = 30
    for c in 'BCDEFG':
        ws1.column_dimensions[c].width = 15

    # Proveedores sheet
    ws2 = wb.create_sheet('Proveedores')
    headers_pr = ['Razón Social', 'NIT', 'Categoría', 'Contacto', 'Teléfono', 'Email', 'Estado']
    for col, h in enumerate(headers_pr, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.fill = amber_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
    for i, p in enumerate(Proveedor.objects.all(), 2):
        vals = [p.razon_social, p.nit, p.categoria, p.contacto, p.telefono, p.email, 'Activo' if p.estado == 'activo' else 'Inactivo']
        for col, v in enumerate(vals, 1):
            c = ws2.cell(row=i, column=col, value=v)
            c.font = data_font
            c.border = thin_border
    ws2.column_dimensions['A'].width = 30
    for c in 'BCDEFG':
        ws2.column_dimensions[c].width = 18

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=reporte_completo.xlsx'
    wb.save(response)
    return response
